# -*- coding: utf-8 -*-
"""
現場マスタの初期データを、実データ（べんり君）から起こす。

🔴 なぜこれが要るのか
  ShiftMax の勤務マスタには「現場」が入っていない。
  1,593件の正体は〈顧客603種 × 勤務区分〉であり、
  実際の現場名（「◯◯（巡回）」等）はどこにも登録されていない。
  管制は毎日それを手で打っている（7月は延べ 1,899回）。
  → 現場名が残っているのは**日次シートの手入力テキストだけ**。
     ここから起こす以外に、現場マスタを作る方法が無い。
  詳細は docs/data-gap-20260917.md §2。

🔴 現場は〈現場名 × 得意先〉で1件（2026-10-02）
  以前は〈警備先番号 × 現場名〉で数えていたため、日勤／夜A／夜B のある現場が
  区分ごとに別の現場として出ていた（7月で 25件）。
  警備先番号は現場に持たせず、引き渡しのときに〈得意先 × 区分〉で引く
  （supabase/migrations/20261002000000_duty_codes.sql）。
  同じ現場名でも得意先が違えば別の現場（駅名の現場に別の元請けが入る）。

🔴 これは「移行ツール」ではなく「一度きりの初期データ作成」
  scripts/dummy-data.mts と同じく**開発ツールであってシステムではない**。
  src/ の外に置く。第1弾の稼働後は使わない。

🔴 出力は Git に入れない（local/ は .gitignore 済み）
  取引先実名・現場名が入るため、受領資料と同じ扱いにする。

使い方:
    python scripts/extract-sites.py                       # 7月分（受領済みの1本）
    python scripts/extract-sites.py 8月.xlsm 9月.xlsm …   # 何本でも。月をまたいで1つにまとめる
    → local/sites-extracted.csv  確認用（Excel でそのまま開ける）
    → local/sites-import.csv     現場一覧（/masters/import にそのまま通せる）
    → local/duties-import.csv    勤務マスター（同上。現場一覧より先に取り込む）

    取り込む順番：得意先マスター → 勤務マスター → 現場一覧
"""
import collections
import csv
import io
import os
import re
import sys
import unicodedata
import warnings

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl が要る:  pip install openpyxl")

# べんり君のファイルは見出し・入力規則の拡張を含み、openpyxl が毎回警告を出す。読み取りには影響しない
warnings.filterwarnings("ignore", module="openpyxl")

DEFAULT_SRC = os.path.join("docs", "管制_別紙", "べんり君_別紙Ⓑ.xlsm")
OUT = os.path.join("local", "sites-extracted.csv")          # 確認用
OUT_IMPORT = os.path.join("local", "sites-import.csv")        # 取込用（現場一覧）
OUT_DUTIES = os.path.join("local", "duties-import.csv")       # 取込用（勤務マスター）

# 日ごとのシート名：「7.1(チェック済み)」「10.15」など。月.日 で始まるものを拾う
DAY_SHEET = re.compile(r"^\s*(\d{1,2})\.(\d{1,2})(?!\d)")

# 日次シートの列（docs/shiftmax-api-analysis.md §8-2）
# D=請求番号 E=請求名 F=担当名 G=警備先番号 H=現場略称 I=現場名
# J=人数 K/L=開始 M/N=終了 O=昼休憩 P=班名 Q=備考
C_GUARD_NO, C_SHORT, C_NAME = 6, 7, 8
C_HEAD, C_SH, C_SM, C_EH, C_EM, C_BREAK, C_NOTE = 9, 10, 11, 12, 13, 14, 16

# 勤務マスタの列（同 §7-2）B〜Y
M_SITE_CODE, M_GUARD_NO, M_KIND = 1, 2, 4
M_CUST_CODE, M_CUST_NO, M_CUST_NAME, M_CUST_STAFF, M_BILL_NO = 14, 15, 16, 19, 20
M_JURIS, M_JURIS_NAME, M_DEPT_CODE, M_DEPT_NAME = 21, 22, 23, 24

# 🔴 現場ではない区分（勤務マスタに擬似現場として実在する）。
#   これを sites に入れると配置ボードの現場一覧に「有休」が並ぶ。
#   新システムでは assignments.kind='off' 側で持つ（data-gap-20260917.md §3-1）。
NOT_A_SITE = {
    "内勤", "会社欠", "有休.", "有休", "有給", "教育", "新任", "現任",
    "育休", "介護休", "忌引休", "業災休", "産休", "振休", "空白",
    # 🔴 現中（現着中止）は「現場」ではなく配置の状態。
    #   work_kind の dayCancel / nightCancel 側で持つ（requirements.md §8-7 ①は未決）。
    "日勤現中", "夜勤現中",
}


def cell(v):
    return str(v if v is not None else "").strip()


def norm(s):
    """現場名の表記を揃える（全角・半角／前後・連続の空白）。突き合わせにだけ使う。"""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", s)).strip()


def mode_of(counter):
    """最頻値と「何通りあったか」を返す。"""
    if not counter:
        return "", 0
    return counter.most_common(1)[0][0], len(counter)


def read_master(wb):
    """勤務マスタ（警備先番号 → 顧客・区分・管轄）。"""
    master = collections.OrderedDict()
    ws = wb["勤務マスター"]
    for r in ws.iter_rows(min_row=3, max_row=5000, max_col=26, values_only=True):
        no = cell(r[M_GUARD_NO])
        if not no:
            continue
        master[no] = {
            "site_code": cell(r[M_SITE_CODE]),
            "kind": cell(r[M_KIND]),
            "cust_code": cell(r[M_CUST_CODE]),
            "cust_no": cell(r[M_CUST_NO]),
            "cust_name": cell(r[M_CUST_NAME]),
            "cust_staff": cell(r[M_CUST_STAFF]),
            "bill_no": cell(r[M_BILL_NO]),
            "juris": cell(r[M_JURIS]),
            "juris_name": cell(r[M_JURIS_NAME]),
            "dept_code": cell(r[M_DEPT_CODE]),
            "dept_name": cell(r[M_DEPT_NAME]),
        }
    return master


def main():
    srcs = sys.argv[1:] or [DEFAULT_SRC]
    for src in srcs:
        if not os.path.exists(src):
            sys.exit("ファイルが見つからない: %s" % src)

    # ---- 1. 各ファイルの日次シートから行を集める
    #   勤務マスタは**最後に渡したファイル**のものを使う（新しい月ほど後ろに渡す前提）。
    #   🔴 古い月の警備先番号が最新のマスタに無ければ、その行は「顧客を引けなかった」に数える。
    master = None
    raw = []  # (日付キー, 警備先番号, 区分, 現場名, 行)
    print("■ 読んだファイル")
    for src in srcs:
        wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
        sheets = rows = 0
        for name in wb.sheetnames:
            m = DAY_SHEET.match(name)
            if not m:
                continue
            sheets += 1
            day = (int(m.group(1)), int(m.group(2)))
            ws = wb[name]
            for r in ws.iter_rows(min_row=9, max_row=400, max_col=21, values_only=True):
                no, nm = cell(r[C_GUARD_NO]), cell(r[C_NAME])
                if not no or not nm:
                    continue
                rows += 1
                raw.append((day, no, cell(r[C_SHORT]), nm, r))
        if "勤務マスター" in wb.sheetnames:
            master = read_master(wb)
        wb.close()
        # 🔴 0件は疑う（シート名の形が違うと、黙って何も拾わずに終わる）
        print("  %s … 日次シート %d 枚 / %d 行%s" % (
            src, sheets, rows, "  🔴 0件。シート名の形を確かめること" if rows == 0 else ""))
    if master is None:
        sys.exit("勤務マスター のシートがどのファイルにも無い")

    # ---- 2. 〈現場名 × 得意先〉でまとめる
    site = collections.OrderedDict()
    skipped = unmatched = 0
    for day, no, kind, nm, r in raw:
        if kind in NOT_A_SITE:      # 休み・内勤などは現場ではない
            skipped += 1
            continue
        mrow = master.get(no)
        if mrow is None:
            unmatched += 1
        staff = mrow["cust_staff"] if mrow else ""
        key = (norm(nm), staff)
        s = site.setdefault(key, {
            "names": collections.Counter(), "nos": collections.Counter(),
            "days": set(), "kind": collections.Counter(),
            "time": collections.Counter(), "brk": collections.Counter(),
            "head": collections.Counter(), "note": collections.Counter(),
        })
        s["names"][nm] += 1
        s["nos"][no] += 1
        s["days"].add(day)
        s["kind"][kind] += 1
        if r[C_SH] is not None and r[C_EH] is not None:
            s["time"]["%02d:%02d-%02d:%02d" % (
                int(r[C_SH]), int(r[C_SM] or 0), int(r[C_EH]), int(r[C_EM] or 0))] += 1
        if r[C_BREAK] is not None:
            s["brk"][str(r[C_BREAK])] += 1
        if r[C_HEAD] is not None:
            try:
                s["head"][int(r[C_HEAD])] += 1
            except (TypeError, ValueError):
                pass
        if r[C_NOTE]:
            s["note"][cell(r[C_NOTE])] += 1

    # ---- 3. 書き出す
    #   ① 確認用   … 人数・出た日数・ばらつきが分かる形
    #   ② 現場一覧 … /masters/import にそのまま通す形。現場コードは持たない（DB が振る）
    #   ③ 勤務マスター … 同上。非表示シートを人が CSV にする手間を省く
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    rows = 0
    merged = 0
    blank = {k: "" for k in ("site_code", "kind", "cust_code", "cust_no", "cust_name",
                             "cust_staff", "bill_no", "juris", "juris_name",
                             "dept_code", "dept_name")}
    # utf-8-sig ＝ Excel が文字化けせずに開ける
    with io.open(OUT, "w", encoding="utf-8-sig", newline="") as f, \
            io.open(OUT_IMPORT, "w", encoding="utf-8-sig", newline="") as g:
        w = csv.writer(f)
        imp = csv.writer(g)
        w.writerow([
            "No", "現場名", "区分", "出た日数", "警備先番号",
            "開始", "終了", "時間は一定", "休憩", "人数", "人数は一定",
            "顧客名", "管轄", "備考に出た語", "確認",
        ])
        imp.writerow([
            "現場", "略称", "フリガナ", "現場住所",
            "予定開始時間", "予定開始分", "予定終了時間", "予定終了分", "予定休憩時間",
            "勤務予定フラグ", "顧客コード", "得意先番号", "担当コード", "請求番号",
            "現場管轄", "管轄表示", "部署コード", "部署名",
        ])
        for (_, staff), s in sorted(site.items(), key=lambda kv: -len(kv[1]["days"])):
            nm, _ = mode_of(s["names"])       # 表記が揺れていたら、いちばん多い書き方を採る
            no, _ = mode_of(s["nos"])         # 管轄・請求番号は最も多く出た番号から引く
            if len(s["nos"]) > 1:
                merged += 1
            m = master.get(no) or blank
            tm, tn = mode_of(s["time"])
            hd, hn = mode_of(s["head"])
            bk, _ = mode_of(s["brk"])
            start, end = (tm.split("-") + [""])[:2] if tm else ("", "")
            rows += 1
            w.writerow([
                rows, nm, "/".join(k for k, _ in s["kind"].most_common()), len(s["days"]),
                "/".join(s["nos"]),
                start, end, "○" if tn == 1 else "△(%d通り)" % tn,
                bk, hd, "○" if hn == 1 else "△(%d通り)" % hn,
                m["cust_name"], m["juris"],
                " / ".join(k for k, _ in s["note"].most_common(2)), "",
            ])

            # 🔴 時刻は「毎回同じ」現場にだけ入れる。
            #   ばらつく現場に代表値を入れると、管制が直さない限り嘘の予定が残る。
            #   空にしておけば取込側は null（＝予定を持たない）として扱う。
            fixed = (tn == 1)
            sh = sm = eh = em = ""
            if fixed and tm:
                sh, sm = start.split(":")
                eh, em = end.split(":")
            imp.writerow([
                nm, nm[:8], "", "",
                sh, sm, eh, em, (bk if fixed else ""),
                "1" if fixed else "0",
                m["cust_code"], m["cust_no"], staff, m["bill_no"],
                m["juris"], m["juris_name"], m["dept_code"], m["dept_name"],
            ])

    with io.open(OUT_DUTIES, "w", encoding="utf-8-sig", newline="") as h:
        d = csv.writer(h)
        d.writerow(["現場コード", "警備先番号", "略称", "担当コード",
                    "顧客コード", "得意先番号", "請求番号"])
        for no, m in master.items():
            d.writerow([m["site_code"], no, m["kind"], m["cust_staff"],
                        m["cust_code"], m["cust_no"], m["bill_no"]])

    print("■ 現場として書き出した: %d 件（うち区分の違う番号をまとめたもの %d 件）" % (rows, merged))
    print("■ 現場ではないため除外した行（休み・内勤など）: %d 行" % skipped)
    print("■ 勤務マスタに警備先番号が無く顧客を引けなかった行: %d 行" % unmatched)
    print("■ 勤務マスター: %d 件" % len(master))
    print("→ 確認用: %s" % OUT)
    print("→ 取込用: %s / %s（/masters/import に通せる）" % (OUT_DUTIES, OUT_IMPORT))


if __name__ == "__main__":
    main()
