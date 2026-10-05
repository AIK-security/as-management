# -*- coding: utf-8 -*-
"""
列車見張の資格（鉄道会社ごと）を、受領した一覧から SQL にする（2026-10-05）。

🔴 読むのは「元データ_リスト」シートの 資格／失効／氏名／認定証番号／交付年月日／有効期限 だけ。
  同じシートに **生年月日・電話番号・ID・PW（LINGS のログイン情報）** があるが、読まない。
  受領資料と同じ扱い（docs/管制_別紙/ に置き、Git に入れない）。

🔴 隊員は氏名で引く（一覧に個人コードが無いため）。
  社員マスターは最新のべんり君（引数の2つ目）から読み、同名が2人いる氏名は使わない（取り違え防止）。

🔴 入れないもの
  ・失効の列に印（△ など）がある行
  ・社員マスターに居ない氏名（辞めた人・表記ゆれ）── 一覧に出すので、要るなら画面から足す

🔴 これは一度きりの取込用（extract-sites.py と同じ開発ツール）。
  出力は Git に入れない（local/ は .gitignore 済み）。

使い方:
    python scripts/extract-train-quals.py <列車資格の一覧.xlsx> <最新のべんり君.xlsm>
    → local/train-quals.sql   SQL Editor で流す
      （先に supabase/migrations/20261005120000_train_qualifications.sql を流しておく）
"""
import importlib.util
import os
import sys
from datetime import date, datetime, timedelta

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("extract_assignments", os.path.join(HERE, "extract-assignments.py"))
ea = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(ea)

OUT = os.path.join("local", "train-quals.sql")
SHEET = "元データ_リスト"

# 一覧の「資格」の文字 → 資格マスタの code（20261005120000_train_qualifications.sql）
CODE_OF = {
    "東急電鉄": "q-train-tokyu",
    "京王電鉄": "q-train-keio",
    "小田急電鉄": "q-train-odakyu",
    "京急電鉄": "q-train-keikyu",
    "相模鉄道": "q-train-sotetsu",
    "JR列車": "q-train-jr",
    "江ノ島電鉄": "q-train-enoden",
    "横浜高速鉄道": "q-train-yokohama",
    "飛島建設": "q-train-tobishima",
}


# 一覧と社員マスターで字が違う人（柴山に同一人物と確認済み・2026-10-05）。一覧の表記 → マスタの表記
ALIAS = {
    "小出佑太": "小出祐太",
    "田中翔太": "田中翔大",
}


def to_date(v) -> date | None:
    """日付のセル。Excel のシリアル値（数値）で入っている行もある。"""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and v > 0:
        return date(1899, 12, 30) + timedelta(days=int(v))
    return None


def main() -> None:
    if len(sys.argv) != 3:
        sys.exit("使い方: python scripts/extract-train-quals.py <列車資格の一覧.xlsx> <最新のべんり君.xlsm>")
    src, benri = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(benri, data_only=True, read_only=True)
    by_name, _ = ea.read_guards(wb)
    wb.close()

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    rows = list(wb[SHEET].iter_rows(values_only=True))
    wb.close()
    head_i = next(i for i, r in enumerate(rows) if r and "氏名" in [str(c).strip() if c else "" for c in r])
    head = [str(c).strip() if c else "" for c in rows[head_i]]
    c_qual, c_lost, c_name = head.index("資格"), head.index("失効"), head.index("氏名")
    c_no, c_issued, c_exp = head.index("認定証番号"), head.index("交付年月日"), head.index("有効期限")

    values = []
    lost, no_guard, unknown_qual = [], set(), set()
    for r in rows[head_i + 1:]:
        if not r[c_name]:
            continue
        name = ea.nz(str(r[c_name]))
        name = ALIAS.get(name, name)
        qual = str(r[c_qual] or "").strip()
        if r[c_lost] not in (None, ""):
            lost.append((name, qual))
            continue
        code = CODE_OF.get(qual)
        if code is None:
            unknown_qual.add(qual)
            continue
        staff = by_name.get(name)
        if staff is None:
            no_guard.add(name)
            continue
        number = r[c_no]
        values.append("(%s, %s, %s, %s, %s)" % (
            ea.q(staff), ea.q(code),
            ea.q(str(number).strip()) if number not in (None, "") else "null",
            ea.q(to_date(r[c_issued]).isoformat()) if to_date(r[c_issued]) else "null",
            ea.q(to_date(r[c_exp]).isoformat()) if to_date(r[c_exp]) else "null",
        ))

    if unknown_qual:
        sys.exit("止めました：資格マスタに対応の無い資格名 %s" % sorted(unknown_qual))

    os.makedirs("local", exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("-- 列車見張の資格（受領一覧から生成）。🔴 Git に入れない。実名が入る\n")
        f.write("-- 生成: scripts/extract-train-quals.py / %d 件\n" % len(values))
        f.write("begin;\n")
        f.write("create temp table imp_tq (staff_code text, code text, number text, issued_on date, expires_on date) on commit drop;\n")
        f.write("insert into imp_tq values\n" + ",\n".join(values) + ";\n")
        f.write("""
-- 隊員・資格を引けない行があれば止める（黙って落とさない）
do $$
declare n int;
begin
  select count(*) into n from imp_tq i
   where not exists (select 1 from public.guards g where g.staff_code = i.staff_code)
      or not exists (select 1 from public.qualifications q where q.code = i.code);
  if n > 0 then raise exception '止めました：隊員か資格を引けない行が % 件', n; end if;
end $$;

-- 🔴 同じ隊員×資格が既にあれば、期限などを一覧の値で上書きする（一覧が新しい）
insert into public.guard_qualifications (guard_id, qualification_id, number, issued_on, expires_on)
select g.id, q.id, i.number, i.issued_on, i.expires_on
  from imp_tq i
  join public.guards g on g.staff_code = i.staff_code
  join public.qualifications q on q.code = i.code
on conflict (guard_id, qualification_id)
do update set number = excluded.number, issued_on = excluded.issued_on, expires_on = excluded.expires_on;

-- 確認用
select (select count(*) from imp_tq) as 一覧の件数,
       (select count(*) from public.guard_qualifications gq
          join public.qualifications q on q.id = gq.qualification_id
         where q.category = 'train') as 列車見張の登録数;
commit;
""")

    print("■ 取り込む: %d 件" % len(values))
    print("■ 失効の印があり入れない: %d 件 %s" % (len(lost), lost))
    print("■ 社員マスターに居ない氏名（入れない）: %d 名 %s" % (len(no_guard), sorted(no_guard)))
    print("→ " + OUT)


if __name__ == "__main__":
    main()
